from parser.base import BaseParser
from parser.models import ParsedDocument, ParsedSection
from utils.md5 import get_string_md5
import docx


class DOCXParser(BaseParser):
    """DOCX 文件解析器"""

    def parse(self, file) -> ParsedDocument:
        doc = docx.Document(file)
        doc_id = get_string_md5(getattr(file, "name", str(file)))
        filename = getattr(file, "name", str(file))
        total_pages = len(doc.paragraphs)
        sections: list[ParsedSection] = []
        idx = 0
        for p in doc.paragraphs:
            if not p.text.strip():
                continue
            idx += 1
            sections.append(
                ParsedSection(title=None, content=p.text, page=idx, level=0)
            )
        return ParsedDocument(doc_id, filename, total_pages, sections)
